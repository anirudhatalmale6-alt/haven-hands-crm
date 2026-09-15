"""Excel / CSV export of a Haven Hands salary payment schedule."""

from __future__ import annotations

import csv
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .salary import SchedulePeriod, SalaryTerms, schedule_totals

COLUMNS = [
    ("number", "#", 6),
    ("period_start", "Period from", 13),
    ("payment_date", "Date of Payment", 15),
    ("basic_salary", "Basic Salary ($)", 15),
    ("rest_days_in_period", "Rest days in period", 11),
    ("off_days_compensated", "Off days compensated", 11),
    ("off_day_comp", "Off Day Comp ($)", 15),
    ("placement_fee_deduction", "Placement Fee ($)", 16),
    ("amount_received", "Amount Received ($)", 18),
    ("loan_balance_after", "Balance of Fee ($)", 16),
    ("remarks", "Remarks", 38),
]

MONEY_COLS = {"basic_salary", "off_day_comp", "placement_fee_deduction",
              "amount_received", "loan_balance_after"}

NAVY = "1F3A5F"
LIGHT = "EEF2F7"


def schedule_to_xlsx(rows: list[SchedulePeriod], terms: SalaryTerms,
                     employer_name: str, fdw_name: str, path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Salary schedule"

    thin = Side(style="thin", color="C8D2DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "SALARY PAYMENT SCHEDULE"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    ws["A2"] = "HAVEN HANDS PTE. LTD.  |  EA Licence No. 25C3257"
    ws["A2"].font = Font(size=10, color="64748B")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNS))

    meta = [
        ("Name of Employer", employer_name),
        ("Name of FDW", fdw_name),
        ("Total Placement Fees", f"SGD {terms.total_placement_fee:,.2f}"),
        ("Date of Commencement", terms.commencement_date.strftime("%d-%m-%Y")),
        ("Basic Salary", f"SGD {terms.basic_salary:,.2f}"),
        ("Salary Per Day", f"SGD {terms.daily_rate:,.2f}"),
        ("Day Off Per Month", terms.days_off_per_month),
        ("Loan Months", terms.loan_months),
    ]
    r = 4
    for label, value in meta:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=3, value=value)
        r += 1

    head = r + 1
    for i, (_key, title, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=head, column=i, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[c.column_letter].width = width

    for n, row in enumerate(rows):
        excel_row = head + 1 + n
        for i, (key, _title, _w) in enumerate(COLUMNS, start=1):
            value = getattr(row, key)
            if isinstance(value, Decimal):
                value = float(value)
            elif hasattr(value, "isoformat"):
                value = value.strftime("%d-%m-%Y")
            c = ws.cell(row=excel_row, column=i, value=value)
            c.border = border
            if key in MONEY_COLS:
                c.number_format = "#,##0.00"
            if n % 2 == 0:
                c.fill = PatternFill("solid", fgColor=LIGHT)
            if row.is_final_instalment and key in ("placement_fee_deduction", "remarks"):
                c.font = Font(bold=True, color="A33A00")

    totals = schedule_totals(rows)
    t_row = head + 1 + len(rows)
    ws.cell(row=t_row, column=1, value="Total").font = Font(bold=True)
    for key, col in (("total_basic", 4), ("total_off_day_comp", 7),
                     ("total_placement_fee", 8), ("total_received", 9)):
        c = ws.cell(row=t_row, column=col, value=float(totals[key]))
        c.font = Font(bold=True)
        c.number_format = "#,##0.00"
        c.border = border

    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    wb.save(path)
    return path


def schedule_to_csv(rows: list[SchedulePeriod]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([title for _k, title, _w in COLUMNS])
    for row in rows:
        w.writerow([getattr(row, key) for key, _t, _w in COLUMNS])
    return buf.getvalue()
